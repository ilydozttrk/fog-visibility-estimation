"""
train_vgg16_attention.py

Training pipeline for the CBAM-enhanced VGG16 regression model
used in fog visibility estimation.

The training configuration is intentionally kept consistent with
the VGG16 baseline so that the effect of the attention mechanism
can be evaluated under comparable experimental conditions.

Author: Ilayda Ozturk
Project: TUBITAK 2209-A
"""

from datetime import datetime
from pathlib import Path

import matplotlib.pyplot as plt
import torch
import torch.nn as nn
from openpyxl import Workbook, load_workbook
from torch.optim import Adam
from torch.utils.data import DataLoader

from src.models.vgg16_attention import (
    build_vgg16_attention_regression_model,
)
from src.training.config import (
    BATCH_SIZE,
    CHECKPOINTS_DIR,
    DEVICE,
    DROPOUT_RATE,
    FREEZE_BACKBONE,
    IMAGE_SIZE,
    LEARNING_RATE,
    NUM_EPOCHS,
    NUM_WORKERS,
    RANDOM_SEED,
    REGRESSION_HIDDEN_DIM_1,
    REGRESSION_HIDDEN_DIM_2,
    TEST_RATIO,
    TRAIN_RATIO,
    VAL_RATIO,
    VGG16_ATTENTION_CHECKPOINT_PATH,
    WEIGHT_DECAY,
    create_output_directories,
)
from src.training.dataloader import create_dataloaders
from src.training.utils import set_seed


# =============================================================================
# Attention Configuration
# =============================================================================

CBAM_REDUCTION_RATIO = 16
CBAM_SPATIAL_KERNEL_SIZE = 7


# =============================================================================
# Training
# =============================================================================

def train_one_epoch(
    model: nn.Module,
    dataloader: DataLoader,
    loss_function: nn.Module,
    optimizer: torch.optim.Optimizer,
    device: torch.device,
) -> float:
    """
    Train the attention-enhanced model for one epoch
    and return mean training MAE.
    """

    model.train()

    running_loss = 0.0
    total_samples = 0

    for images, targets in dataloader:
        images = images.to(device)
        targets = targets.to(device)

        optimizer.zero_grad()

        predictions = model(images).squeeze(1)

        loss = loss_function(
            predictions,
            targets,
        )

        loss.backward()
        optimizer.step()

        batch_size = images.size(0)

        running_loss += loss.item() * batch_size
        total_samples += batch_size

    mean_loss = running_loss / total_samples

    return mean_loss


# =============================================================================
# Validation
# =============================================================================

def validate_one_epoch(
    model: nn.Module,
    dataloader: DataLoader,
    loss_function: nn.Module,
    device: torch.device,
) -> float:
    """
    Evaluate the model on the validation dataset
    and return mean validation MAE.
    """

    model.eval()

    running_loss = 0.0
    total_samples = 0

    with torch.no_grad():
        for images, targets in dataloader:
            images = images.to(device)
            targets = targets.to(device)

            predictions = model(images).squeeze(1)

            loss = loss_function(
                predictions,
                targets,
            )

            batch_size = images.size(0)

            running_loss += loss.item() * batch_size
            total_samples += batch_size

    mean_loss = running_loss / total_samples

    return mean_loss


# =============================================================================
# Checkpoint
# =============================================================================

def save_checkpoint(
    model: nn.Module,
    optimizer: torch.optim.Optimizer,
    epoch: int,
    validation_mae: float,
    checkpoint_path: Path,
) -> None:
    """
    Save the best attention model checkpoint
    and related training state.
    """

    checkpoint = {
        "epoch": epoch,
        "model_state_dict": model.state_dict(),
        "optimizer_state_dict": optimizer.state_dict(),
        "validation_mae": validation_mae,
        "random_seed": RANDOM_SEED,
        "model_name": "VGG16_CBAM",
        "cbam_reduction_ratio": CBAM_REDUCTION_RATIO,
        "cbam_spatial_kernel_size": CBAM_SPATIAL_KERNEL_SIZE,
    }

    checkpoint_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    torch.save(
        checkpoint,
        checkpoint_path,
    )


# =============================================================================
# Training Outputs
# =============================================================================

def save_training_curve(
    training_mae_history: list[float],
    validation_mae_history: list[float],
    output_path: Path,
) -> None:
    """
    Plot and save training and validation MAE values by epoch.
    """

    epochs = range(
        1,
        len(training_mae_history) + 1,
    )

    plt.figure(
        figsize=(10, 6)
    )

    plt.plot(
        epochs,
        training_mae_history,
        marker="o",
        label="Training MAE",
    )

    plt.plot(
        epochs,
        validation_mae_history,
        marker="o",
        label="Validation MAE",
    )

    plt.xlabel("Epoch")
    plt.ylabel("MAE (m)")
    plt.title(
        "VGG16 + CBAM Learning Curve"
    )

    plt.legend()
    plt.grid(
        True,
        alpha=0.3,
    )

    plt.tight_layout()

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    plt.savefig(
        output_path,
        dpi=300,
        bbox_inches="tight",
    )

    plt.close()


def record_experiment(
    experiment_log_path: Path,
    experiment_id: str,
    training_mae_history: list[float],
    validation_mae_history: list[float],
    best_epoch: int,
    best_validation_mae: float,
) -> None:
    """
    Append the completed attention experiment
    to experiment_log.xlsx.
    """

    headers = [
        "Experiment ID",
        "Date",
        "Model",
        "Pretrained Weights",
        "Backbone Frozen",
        "Image Size",
        "Batch Size",
        "Epochs",
        "Optimizer",
        "Learning Rate",
        "Weight Decay",
        "Loss Function",
        "Random Seed",
        "Final Train MAE",
        "Final Validation MAE",
        "Best Validation MAE",
        "Best Epoch",
        "Checkpoint Path",
        "Notes",
    ]

    if experiment_log_path.exists():
        workbook = load_workbook(
            experiment_log_path
        )

        worksheet = workbook.active

    else:
        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = "Experiments"

        worksheet.append(headers)

    worksheet.append(
        [
            experiment_id,
            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            ),
            "VGG16 + CBAM",
            "ImageNet",
            FREEZE_BACKBONE,
            f"{IMAGE_SIZE}x{IMAGE_SIZE}",
            BATCH_SIZE,
            NUM_EPOCHS,
            "Adam",
            LEARNING_RATE,
            WEIGHT_DECAY,
            "L1Loss / MAE",
            RANDOM_SEED,
            training_mae_history[-1],
            validation_mae_history[-1],
            best_validation_mae,
            best_epoch,
            str(
                VGG16_ATTENTION_CHECKPOINT_PATH
            ),
            (
                "VGG16 baseline enhanced with CBAM. "
                "CBAM inserted before the final VGG16 MaxPool."
            ),
        ]
    )

    for column_cells in worksheet.columns:
        maximum_length = 0

        column_letter = (
            column_cells[0].column_letter
        )

        for cell in column_cells:
            if cell.value is not None:
                maximum_length = max(
                    maximum_length,
                    len(str(cell.value)),
                )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            maximum_length + 2,
            40,
        )

    experiment_log_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook.save(
        experiment_log_path
    )


def save_initial_training_log(
    log_path: Path,
    experiment_id: str,
    training_mae_history: list[float],
    validation_mae_history: list[float],
    best_epoch: int,
    best_validation_mae: float,
) -> None:
    """
    Save a human-readable summary of the
    initial VGG16 + CBAM training run.
    """

    lines = [
        "# Initial VGG16 + CBAM Training Log",
        "",
        f"- Experiment ID: {experiment_id}",
        (
            "- Date: "
            f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ),
        "- Model: VGG16 + CBAM",
        "- Pretrained weights: ImageNet",
        f"- Frozen backbone: {FREEZE_BACKBONE}",
        f"- Device: {DEVICE}",
        (
            f"- Image size: "
            f"{IMAGE_SIZE} x {IMAGE_SIZE}"
        ),
        f"- Batch size: {BATCH_SIZE}",
        f"- Epoch count: {NUM_EPOCHS}",
        "- Optimizer: Adam",
        f"- Learning rate: {LEARNING_RATE}",
        f"- Weight decay: {WEIGHT_DECAY}",
        (
            "- Loss function: "
            "L1Loss / Mean Absolute Error"
        ),
        f"- Random seed: {RANDOM_SEED}",
        (
            "- CBAM reduction ratio: "
            f"{CBAM_REDUCTION_RATIO}"
        ),
        (
            "- CBAM spatial kernel size: "
            f"{CBAM_SPATIAL_KERNEL_SIZE}"
        ),
        (
            "- CBAM integration point: "
            "after VGG16 conv5_3 and before final MaxPool"
        ),
        "",
        "## Epoch Results",
        "",
        (
            "| Epoch | Training MAE (m) | "
            "Validation MAE (m) |"
        ),
        "|---:|---:|---:|",
    ]

    for epoch, (
        training_mae,
        validation_mae,
    ) in enumerate(
        zip(
            training_mae_history,
            validation_mae_history,
        ),
        start=1,
    ):
        lines.append(
            f"| {epoch} | "
            f"{training_mae:.4f} | "
            f"{validation_mae:.4f} |"
        )

    lines.extend(
        [
            "",
            "## Best Result",
            "",
            f"- Best epoch: {best_epoch}",
            (
                "- Best validation MAE: "
                f"{best_validation_mae:.4f} m"
            ),
            (
                "- Checkpoint: "
                f"`{VGG16_ATTENTION_CHECKPOINT_PATH}`"
            ),
            "",
            "## Initial Observation",
            "",
            (
                "The VGG16 + CBAM training pipeline "
                "completed successfully. "
                "The checkpoint corresponding to the "
                "lowest validation MAE was saved for "
                "subsequent independent evaluation and "
                "comparison with the VGG16 baseline."
            ),
        ]
    )

    log_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    log_path.write_text(
        "\n".join(lines),
        encoding="utf-8",
    )


# =============================================================================
# Training Pipeline
# =============================================================================

def train_model() -> None:
    """
    Run the complete VGG16 + CBAM training pipeline.
    """

    set_seed(RANDOM_SEED)
    create_output_directories()

    pin_memory = DEVICE.type == "cuda"

    train_loader, val_loader, _ = (
        create_dataloaders(
            image_size=IMAGE_SIZE,
            batch_size=BATCH_SIZE,
            num_workers=NUM_WORKERS,
            train_ratio=TRAIN_RATIO,
            val_ratio=VAL_RATIO,
            test_ratio=TEST_RATIO,
            random_seed=RANDOM_SEED,
            pin_memory=pin_memory,
        )
    )

    model = (
        build_vgg16_attention_regression_model(
            freeze_backbone=FREEZE_BACKBONE,
            regression_hidden_dim_1=(
                REGRESSION_HIDDEN_DIM_1
            ),
            regression_hidden_dim_2=(
                REGRESSION_HIDDEN_DIM_2
            ),
            dropout_rate=DROPOUT_RATE,
            cbam_reduction_ratio=(
                CBAM_REDUCTION_RATIO
            ),
            cbam_spatial_kernel_size=(
                CBAM_SPATIAL_KERNEL_SIZE
            ),
        )
    )

    model = model.to(DEVICE)

    loss_function = nn.L1Loss()

    trainable_parameters = [
        parameter
        for parameter in model.parameters()
        if parameter.requires_grad
    ]

    optimizer = Adam(
        trainable_parameters,
        lr=LEARNING_RATE,
        weight_decay=WEIGHT_DECAY,
    )

    best_validation_mae = float("inf")
    best_epoch = 0

    training_mae_history: list[float] = []
    validation_mae_history: list[float] = []

    results_dir = CHECKPOINTS_DIR.parent
    plots_dir = results_dir / "plots"
    logs_dir = results_dir / "logs"

    loss_curve_path = (
        plots_dir
        / "vgg16_attention_loss_curve.png"
    )

    experiment_log_path = (
        results_dir
        / "experiment_log.xlsx"
    )

    initial_training_log_path = (
        logs_dir
        / "vgg16_attention_initial_training_log.md"
    )

    experiment_id = datetime.now().strftime(
        "VGG16_CBAM_%Y%m%d_%H%M%S"
    )

    total_parameters = sum(
        parameter.numel()
        for parameter in model.parameters()
    )

    trainable_parameter_count = sum(
        parameter.numel()
        for parameter in model.parameters()
        if parameter.requires_grad
    )

    print("=" * 70)
    print(
        "VGG16 + CBAM Fog Visibility Regression Training"
    )
    print("=" * 70)

    print(
        f"Device                : {DEVICE}"
    )

    print(
        "Training images       : "
        f"{len(train_loader.dataset)}"
    )

    print(
        "Validation images     : "
        f"{len(val_loader.dataset)}"
    )

    print(
        f"Batch size            : {BATCH_SIZE}"
    )

    print(
        f"Epochs                : {NUM_EPOCHS}"
    )

    print(
        f"Learning rate         : {LEARNING_RATE}"
    )

    print(
        f"Random seed           : {RANDOM_SEED}"
    )

    print(
        f"Frozen backbone       : {FREEZE_BACKBONE}"
    )

    print(
        "CBAM reduction ratio  : "
        f"{CBAM_REDUCTION_RATIO}"
    )

    print(
        "CBAM spatial kernel   : "
        f"{CBAM_SPATIAL_KERNEL_SIZE}"
    )

    print(
        "Total parameters      : "
        f"{total_parameters:,}"
    )

    print(
        "Trainable parameters  : "
        f"{trainable_parameter_count:,}"
    )

    print(
        "Checkpoint path       : "
        f"{VGG16_ATTENTION_CHECKPOINT_PATH}"
    )

    print("=" * 70)

    for epoch in range(
        1,
        NUM_EPOCHS + 1,
    ):
        training_mae = train_one_epoch(
            model=model,
            dataloader=train_loader,
            loss_function=loss_function,
            optimizer=optimizer,
            device=DEVICE,
        )

        validation_mae = validate_one_epoch(
            model=model,
            dataloader=val_loader,
            loss_function=loss_function,
            device=DEVICE,
        )

        training_mae_history.append(
            training_mae
        )

        validation_mae_history.append(
            validation_mae
        )

        print(
            f"Epoch [{epoch:02d}/{NUM_EPOCHS}] "
            f"| Train MAE: {training_mae:.4f} m "
            f"| Validation MAE: "
            f"{validation_mae:.4f} m"
        )

        if (
            validation_mae
            < best_validation_mae
        ):
            best_validation_mae = (
                validation_mae
            )

            best_epoch = epoch

            save_checkpoint(
                model=model,
                optimizer=optimizer,
                epoch=epoch,
                validation_mae=validation_mae,
                checkpoint_path=(
                    VGG16_ATTENTION_CHECKPOINT_PATH
                ),
            )

            print(
                "  -> Best checkpoint saved "
                "(Validation MAE: "
                f"{best_validation_mae:.4f} m)"
            )

    save_training_curve(
        training_mae_history=(
            training_mae_history
        ),
        validation_mae_history=(
            validation_mae_history
        ),
        output_path=loss_curve_path,
    )

    record_experiment(
        experiment_log_path=(
            experiment_log_path
        ),
        experiment_id=experiment_id,
        training_mae_history=(
            training_mae_history
        ),
        validation_mae_history=(
            validation_mae_history
        ),
        best_epoch=best_epoch,
        best_validation_mae=(
            best_validation_mae
        ),
    )

    save_initial_training_log(
        log_path=initial_training_log_path,
        experiment_id=experiment_id,
        training_mae_history=(
            training_mae_history
        ),
        validation_mae_history=(
            validation_mae_history
        ),
        best_epoch=best_epoch,
        best_validation_mae=(
            best_validation_mae
        ),
    )

    print("=" * 70)
    print("Training completed.")

    print(
        f"Best epoch         : {best_epoch}"
    )

    print(
        "Best validation MAE: "
        f"{best_validation_mae:.4f} m"
    )

    print(
        "Best model         : "
        f"{VGG16_ATTENTION_CHECKPOINT_PATH}"
    )

    print(
        f"Experiment log     : {experiment_log_path}"
    )

    print(
        f"Loss curve         : {loss_curve_path}"
    )

    print(
        f"Training log       : {initial_training_log_path}"
    )

    print("=" * 70)


def main() -> None:
    train_model()


if __name__ == "__main__":
    main()